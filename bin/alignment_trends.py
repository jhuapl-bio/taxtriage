#!/usr/bin/env python3
# ##############################################################################################
# # Copyright 2025 The Johns Hopkins University Applied Physics Laboratory LLC
# # All rights reserved.
# # Permission is hereby granted, free of charge, to any person obtaining a copy of this
# # software and associated documentation files (the "Software"), to deal in the Software
# # without restriction, including without limitation the rights to use, copy, modify,
# # merge, publish, distribute, sublicense, and/or sell copies of the Software, and to
# # permit persons to whom the Software is furnished to do so.
# #
# # THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR IMPLIED,
# # INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY, FITNESS FOR A PARTICULAR
# # PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE AUTHORS OR COPYRIGHT HOLDERS BE
# # LIABLE FOR ANY CLAIM, DAMAGES OR OTHER LIABILITY, WHETHER IN AN ACTION OF CONTRACT,
# # TORT OR OTHERWISE, ARISING FROM, OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE
# # OR OTHER DEALINGS IN THE SOFTWARE.
# #
"""
alignment_trends.py — cross-sample trends in alignment depth.

Reads the per-sample TaxTriage JSONs (``*.paths.json``) or a combined
``all.odr.json`` and, for every reference (strain key) seen in two or more
samples, lines the samples' ``depth_profile`` windows up on one grid and asks,
window by window: in how many samples is this stretch of the reference
  * ZERO  — no read touches the window,
  * LOW   — depth below ``--low_frac`` x that sample's own mean depth (or an
            absolute ``--low_abs``), or
  * HIGH  — depth above ``--high_frac`` x the sample's mean (or ``--high_abs``)?

Windows where the same state recurs in at least ``--min_freq`` of the
informative samples are merged into regions. Recurrent ZERO/LOW regions point
at parts of the reference assembly that this sample set does not carry
(deletions, divergent / novel loci, or assembly artefacts); recurrent HIGH
regions point at repeats, rRNA operons, mobile elements, plasmid copy number
or contamination-prone loci.

Depth is normalised per sample (window depth / sample mean) so deep and
shallow libraries are comparable. A sample only counts toward a reference's
frequencies when it expects at least ``--min_reads_per_window`` reads per
window (its reads x window / reference length): with fewer, an empty or thin
window is mostly Poisson sampling noise, not a real gap.

Outputs (``-o PREFIX``):
  PREFIX.regions.tsv   recurrent zero / low / high regions
  PREFIX.windows.tsv   per-window frequencies and normalised depth summary
  PREFIX.samples.tsv   per sample x reference depth / class summary
  PREFIX.matrix.tsv    (--matrix) long per-sample x window table
  PREFIX.json          summary, regions and samples (no per-window table) for the report / other tools
  PREFIX.xlsx          (if openpyxl is installed) regions / samples / windows
  PREFIX.plots/        (--plots N, needs matplotlib) one PNG per top reference

Example:
  alignment_trends.py -i alignment/*.paths.json -o all.alignment_trends --plots 10
"""
import argparse
import csv
import glob
import json
import math
import os
import statistics
import sys
from collections import defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from depth_profile import decode_profile, rebin  # noqa: E402

CLASSES = ("zero", "low", "high")


def parse_args(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("-i", "--input", nargs="+", required=True,
                   help="Per-sample *.paths.json and/or combined all.odr.json files (globs ok).")
    p.add_argument("-o", "--output", default="alignment_trends", help="Output prefix.")
    p.add_argument("--min_samples", type=int, default=2,
                   help="Minimum informative samples for a reference to be analysed (default 2).")
    p.add_argument("--min_reads", type=int, default=10,
                   help="Minimum aligned reads for a sample to count for a reference (default 10).")
    p.add_argument("--min_reads_per_window", type=float, default=20.0,
                   help="A sample only counts toward zero/low/high frequencies on a reference when it "
                        "expects at least this many reads per window (its reads x window / reference length). "
                        "Below that, empty or thin windows are mostly Poisson sampling noise (default 20).")
    p.add_argument("--min_mean_depth", type=float, default=0.0,
                   help="Optional extra floor on a sample's mean depth for it to count (default 0 = off).")
    p.add_argument("--low_frac", type=float, default=0.2,
                   help="LOW = window depth < low_frac x sample mean depth (default 0.2).")
    p.add_argument("--high_frac", type=float, default=3.0,
                   help="HIGH = window depth > high_frac x sample mean depth (default 3.0).")
    p.add_argument("--low_abs", type=float, default=None,
                   help="Absolute LOW depth cutoff; overrides --low_frac when set.")
    p.add_argument("--high_abs", type=float, default=None,
                   help="Absolute HIGH depth cutoff; overrides --high_frac when set.")
    p.add_argument("--zero_breadth", type=float, default=0.0,
                   help="Window counts as ZERO when its breadth %% is <= this (default 0 = no read at all).")
    p.add_argument("--min_freq", type=float, default=0.5,
                   help="Fraction of informative samples a window state must recur in (default 0.5).")
    p.add_argument("--min_region_windows", type=int, default=1,
                   help="Minimum consecutive windows for a recurrent region (default 1).")
    p.add_argument("--window", type=int, default=0,
                   help="Force a coarser common window (bp); rounded up to the samples' window ladder.")
    p.add_argument("--samples", nargs="*", default=None, help="Only use these sample names.")
    p.add_argument("--exclude_controls", action="store_true",
                   help="Skip samples whose metadata marks them as controls / in-silico.")
    p.add_argument("--matrix", action="store_true", help="Also write the long per-sample x window matrix.")
    p.add_argument("--plots", type=int, default=0, help="Write PNGs for the top N references (needs matplotlib).")
    return p.parse_args(argv)


# ─────────────────────────────────────────────────────────────── loading ──
def _iter_sample_dicts(paths):
    seen = set()
    for raw in paths:
        files = sorted(glob.glob(raw)) or [raw]
        for f in files:
            if not os.path.isfile(f) or os.path.basename(f).startswith("NO_FILE"):
                continue
            try:
                with open(f) as fh:
                    data = json.load(fh)
            except Exception as exc:
                print(f"[alignment_trends] WARNING: cannot read {f}: {exc}", file=sys.stderr)
                continue
            if isinstance(data, dict) and isinstance(data.get("samples"), list):
                items = data["samples"]
            else:
                items = [data]
            for d in items:
                if not isinstance(d, dict) or "organisms" not in d:
                    continue
                meta = d.get("metadata") or {}
                name = str(meta.get("sample_name") or os.path.basename(f).split(".paths.json")[0])
                if name in seen:
                    print(f"[alignment_trends] WARNING: duplicate sample {name} in {f}; keeping first",
                          file=sys.stderr)
                    continue
                seen.add(name)
                yield name, meta, d


def _iter_strains(d):
    for g in d.get("organisms") or []:
        for sk in g.get("members") or []:
            for st in sk.get("members") or []:
                yield g, sk, st


def load(args):
    """-> refs: {key: {"name", "subkeyname", "toplevelname", "samples": {sample: info}}}"""
    refs = {}
    n_samples = 0
    wanted = set(args.samples) if args.samples else None
    for name, meta, d in _iter_sample_dicts(args.input):
        if wanted and name not in wanted:
            continue
        if args.exclude_controls and (meta.get("insilico") or meta.get("control_type")):
            continue
        n_samples += 1
        for g, sk, st in _iter_strains(d):
            prof = st.get("depth_profile")
            if not prof:
                continue
            key = str(st.get("key", ""))
            ref = refs.setdefault(key, {
                "key": key,
                "name": st.get("name", key),
                "subkeyname": st.get("subkeyname", ""),
                "toplevelname": st.get("toplevelname", g.get("toplevelname", "")),
                "samples": {},
            })
            ref["samples"][name] = {
                "numreads": float(st.get("numreads") or 0),
                "meandepth_json": float(st.get("meandepth") or 0),
                "coverage_json": float(st.get("coverage") or 0),
                "tass": st.get("tass_score"),
                "read_len": float(prof.get("read_len") or st.get("avg_read_length")
                                  or g.get("avg_read_length") or 150),
                "profile_reads": float(prof.get("reads") or 0),
                "profile": prof,
            }
    return refs, n_samples


# ─────────────────────────────────────────────────────────────── analysis ──
def _common_grid(sample_infos, force_window=0):
    """Pick the shared window and the union contig layout."""
    decoded = {}
    wmax = 0
    for s, info in sample_infos.items():
        w, contigs = decode_profile(info["profile"])
        if not contigs:
            continue
        decoded[s] = (w, contigs)
        wmax = max(wmax, w)
    if force_window and force_window > wmax:
        w = wmax or 100
        while w < force_window:
            w *= 2
        wmax = w
    lengths = {}
    for _, (_, contigs) in decoded.items():
        for acc, (length, _, _) in contigs.items():
            lengths[acc] = max(lengths.get(acc, 0), length)
    layout = []  # (acc, length, n_windows)
    for acc in sorted(lengths):
        L = lengths[acc]
        layout.append((acc, L, max(1, (L + wmax - 1) // wmax)))
    return wmax, layout, decoded


def _sample_track(decoded_one, wmax, layout):
    """Rebin one sample onto the common grid -> (depth[], breadth[], winlen[])."""
    w, contigs = decoded_one
    f = max(1, wmax // w)
    dep, br, wl = [], [], []
    for acc, L, n in layout:
        if acc in contigs:
            length, d, b = contigs[acc]
            d2 = rebin(d, f, length, w)
            b2 = rebin(b, f, length, w)
        else:
            d2, b2 = [], []
        d2 = (d2 + [0.0] * n)[:n]
        b2 = (b2 + [0.0] * n)[:n]
        dep.extend(d2)
        br.extend(b2)
        for i in range(n):
            wl.append(max(1, min((i + 1) * wmax, L) - i * wmax))
    return dep, br, wl


def classify(dep, br, mean_depth, args):
    lo = args.low_abs if args.low_abs is not None else args.low_frac * mean_depth
    hi = args.high_abs if args.high_abs is not None else args.high_frac * mean_depth
    out = []
    for d, b in zip(dep, br):
        if b <= args.zero_breadth or d <= 0:
            out.append("zero")
        elif d < lo:
            out.append("low")
        elif d > hi:
            out.append("high")
        else:
            out.append("normal")
    return out


def analyse_reference(ref, args):
    wmax, layout, decoded = _common_grid(ref["samples"], args.window)
    if not layout:
        return None
    # per-window coordinates
    coords = []
    for acc, L, n in layout:
        for i in range(n):
            coords.append((acc, i * wmax, min((i + 1) * wmax, L)))
    nwin = len(coords)
    # Full reference size. Fragmented assemblies only list zero-read contigs up
    # to a cap (bin/depth_profile.py), so contigs no sample covered may be
    # missing from the layout: they are counted here as never-covered bp.
    layout_len = sum(L for _, L, _ in layout)
    full_len = max([layout_len] + [int(i["profile"].get("total_len") or 0) for i in ref["samples"].values()])
    full_contigs = max([len(layout)] + [int(i["profile"].get("n_contigs") or 0) for i in ref["samples"].values()])
    never_bp = max(0, full_len - layout_len)
    never_contigs = max(0, full_contigs - len(layout))

    tracks, sample_rows = {}, []
    for s, info in sorted(ref["samples"].items()):
        if s not in decoded:
            continue
        dep, br, wl = _sample_track(decoded[s], wmax, layout)
        tot = float(full_len)
        mean_d = sum(d * l for d, l in zip(dep, wl)) / tot if tot else 0.0
        breadth = sum(b * l for b, l in zip(br, wl)) / tot if tot else 0.0
        # Reads expected per bp: the sample's reads spread evenly over the
        # reference (falls back to depth / read length for older profiles).
        _reads = info.get("profile_reads") or info["numreads"]
        tot_len = float(full_len)
        per_bp = (_reads / tot_len) if (_reads and tot_len) else \
            mean_d / max(1.0, info.get("read_len") or 150.0)
        exp_reads = per_bp * wmax
        informative = (info["numreads"] >= args.min_reads and mean_d >= args.min_mean_depth
                       and exp_reads >= args.min_reads_per_window)
        cls = classify(dep, br, mean_d, args)
        counts = {c: cls.count(c) for c in CLASSES + ("normal",)}
        # Reads expected in each window at this sample's mean depth: short
        # windows (small contigs, contig ends) need the same evidence as full
        # ones before an empty window is believed.
        win_ok = [informative and per_bp * l >= args.min_reads_per_window for l in wl]
        tracks[s] = dict(dep=dep, br=br, mean=mean_d, cls=cls, informative=informative, win_ok=win_ok)
        sample_rows.append({
            "key": ref["key"], "organism": ref["name"], "sample": s,
            "numreads": int(info["numreads"]), "mean_depth": round(mean_d, 4),
            "breadth_pct": round(breadth, 3), "tass_score": info.get("tass"),
            "exp_reads_per_window": round(exp_reads, 2),
            "informative": informative, "n_windows": nwin,
            "pct_zero": round(100.0 * counts["zero"] / nwin, 2),
            "pct_low": round(100.0 * counts["low"] / nwin, 2),
            "pct_high": round(100.0 * counts["high"] / nwin, 2),
        })

    inf = [s for s, t in tracks.items() if t["informative"]]
    if len(inf) < args.min_samples:
        return {"skipped": True, "samples": sample_rows, "n_informative": len(inf)}

    windows = []
    for wi, (acc, st, en) in enumerate(coords):
        cnt = {c: 0 for c in CLASSES}
        norm = []
        win_inf = [s for s in inf if tracks[s]["win_ok"][wi]]
        for s in win_inf:
            t = tracks[s]
            c = t["cls"][wi]
            if c in cnt:
                cnt[c] += 1
            if t["mean"] > 0:
                norm.append(t["dep"][wi] / t["mean"])
        n = len(win_inf)
        if n < args.min_samples:
            # Too little evidence in this window (e.g. a tiny contig): report
            # it, but it never seeds a recurrent region.
            windows.append({
                "key": ref["key"], "organism": ref["name"], "contig": acc,
                "start": st, "end": en, "n_samples": n,
                "n_zero": cnt["zero"], "n_low": cnt["low"], "n_high": cnt["high"],
                "freq_zero": None, "freq_low": None, "freq_high": None,
                "mean_norm_depth": None, "median_norm_depth": None,
                "sd_log2_norm_depth": None, "mean_log2_norm_depth": None,
            })
            continue
        logs = [math.log2(x + 0.01) for x in norm]
        mu = statistics.mean(logs) if logs else 0.0
        sd = statistics.pstdev(logs) if len(logs) > 1 else 0.0
        windows.append({
            "key": ref["key"], "organism": ref["name"], "contig": acc,
            "start": st, "end": en, "n_samples": n,
            "n_zero": cnt["zero"], "n_low": cnt["low"], "n_high": cnt["high"],
            "freq_zero": round(cnt["zero"] / n, 4),
            "freq_low": round((cnt["low"] + cnt["zero"]) / n, 4),   # low INCLUDES zero
            "freq_high": round(cnt["high"] / n, 4),
            "mean_norm_depth": round(statistics.mean(norm), 4) if norm else 0.0,
            "median_norm_depth": round(statistics.median(norm), 4) if norm else 0.0,
            "sd_log2_norm_depth": round(sd, 4),
            "mean_log2_norm_depth": round(mu, 4),
        })

    regions = []
    for rtype, fkey, member in (("zero", "freq_zero", ("zero",)),
                                ("low", "freq_low", ("zero", "low")),
                                ("high", "freq_high", ("high",))):
        run = []

        def flush():
            if len(run) >= args.min_region_windows:
                ws = [windows[i] for i in run]
                affected = sorted({s for i in run for s in inf
                                   if tracks[s]["win_ok"][i] and tracks[s]["cls"][i] in member})
                regions.append({
                    "key": ref["key"], "organism": ref["name"], "type": rtype,
                    "contig": ws[0]["contig"], "start": ws[0]["start"], "end": ws[-1]["end"],
                    "length": ws[-1]["end"] - ws[0]["start"], "n_windows": len(run),
                    "mean_freq": round(statistics.mean(w[fkey] for w in ws), 4),
                    "max_freq": max(w[fkey] for w in ws),
                    "n_samples": max(w["n_samples"] for w in ws),
                    "mean_norm_depth": round(statistics.mean(w["mean_norm_depth"] for w in ws), 4),
                    "samples_affected": ",".join(affected),
                    "whole_contig": ws[0]["start"] == 0 and
                                    ws[-1]["end"] == next(L for a, L, _ in layout if a == ws[0]["contig"]),
                })
            run.clear()

        prev_contig = None
        for i, w in enumerate(windows):
            if w["contig"] != prev_contig:
                flush()
                prev_contig = w["contig"]
            if w[fkey] is not None and w[fkey] >= args.min_freq:
                run.append(i)
            else:
                flush()
        flush()

    matrix = []
    if args.matrix:
        for s, t in tracks.items():
            for wi, (acc, st, en) in enumerate(coords):
                matrix.append({
                    "key": ref["key"], "organism": ref["name"], "sample": s,
                    "contig": acc, "start": st, "end": en,
                    "depth": round(t["dep"][wi], 3), "breadth_pct": round(t["br"][wi], 1),
                    "norm_depth": round(t["dep"][wi] / t["mean"], 4) if t["mean"] else 0.0,
                    "class": t["cls"][wi], "informative": t["informative"],
                })

    total_len = full_len
    summary = {
        "key": ref["key"], "organism": ref["name"], "subkeyname": ref["subkeyname"],
        "toplevelname": ref["toplevelname"], "window": wmax, "total_len": total_len,
        "n_contigs": full_contigs, "n_windows": nwin,
        "never_covered_contigs": never_contigs, "never_covered_bp": never_bp,
        "n_samples": len(tracks), "n_informative": len(inf),
        "informative_samples": sorted(inf),
    }
    for rtype in ("zero", "low", "high"):
        rs = [r for r in regions if r["type"] == rtype]
        summary[f"n_{rtype}_regions"] = len(rs)
        summary[f"{rtype}_bp"] = sum(r["length"] for r in rs)
        summary[f"{rtype}_pct"] = round(100.0 * summary[f"{rtype}_bp"] / total_len, 3) if total_len else 0
    return {"summary": summary, "windows": windows, "regions": regions,
            "samples": sample_rows, "matrix": matrix, "tracks": tracks,
            "coords": coords, "layout": layout}


# ─────────────────────────────────────────────────────────────── writing ──
def _write_tsv(path, rows, cols=None):
    if not rows and not cols:
        cols = ["note"]
        rows = [{"note": "no data"}]
    cols = cols or list(rows[0].keys())
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


REGION_COLS = ["key", "organism", "type", "contig", "start", "end", "length", "n_windows",
               "mean_freq", "max_freq", "n_samples", "mean_norm_depth", "whole_contig", "samples_affected"]
WINDOW_COLS = ["key", "organism", "contig", "start", "end", "n_samples", "n_zero", "n_low", "n_high",
               "freq_zero", "freq_low", "freq_high", "mean_norm_depth", "median_norm_depth",
               "mean_log2_norm_depth", "sd_log2_norm_depth"]
SAMPLE_COLS = ["key", "organism", "sample", "numreads", "mean_depth", "breadth_pct", "tass_score",
               "exp_reads_per_window", "informative", "n_windows", "pct_zero", "pct_low", "pct_high"]
SUMMARY_COLS = ["key", "organism", "subkeyname", "toplevelname", "window", "total_len", "n_contigs",
                "never_covered_contigs", "never_covered_bp",
                "n_windows", "n_samples", "n_informative", "n_zero_regions", "zero_bp", "zero_pct",
                "n_low_regions", "low_bp", "low_pct", "n_high_regions", "high_bp", "high_pct"]


def _write_xlsx(path, sheets):
    try:
        from openpyxl import Workbook
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for title, cols, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(cols)
        for r in rows[:1_000_000]:
            ws.append([r.get(c) if not isinstance(r.get(c), (list, dict)) else str(r.get(c)) for c in cols])
        ws.freeze_panes = "A2"
    wb.save(path)
    return True


def _plot(res, path):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return False
    tracks = res["tracks"]
    samples = sorted(tracks, key=lambda s: (not tracks[s]["informative"], s))
    nwin = len(res["coords"])
    mat = []
    for s in samples:
        t = tracks[s]
        m = t["mean"] or 1.0
        mat.append([math.log2(d / m + 0.01) if d > 0 else math.nan for d in t["dep"]])
    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(12, 2.2 + 0.28 * len(samples)),
                                   gridspec_kw={"height_ratios": [max(1, len(samples)), 3]}, sharex=True,
                                   constrained_layout=True)
    cmap = plt.get_cmap("RdBu_r").copy()
    cmap.set_bad("#222222")
    im = ax1.imshow(mat, aspect="auto", cmap=cmap, vmin=-3, vmax=3, interpolation="nearest")
    ax1.set_yticks(range(len(samples)))
    ax1.set_yticklabels([s + ("" if tracks[s]["informative"] else " *") for s in samples], fontsize=7)
    # One colorbar spanning both panels keeps their x-axes aligned.
    fig.colorbar(im, ax=[ax1, ax2], label="log2(depth / sample mean); black = zero", pad=0.01, shrink=0.9)
    x = range(nwin)
    w = res["windows"]
    def _f(r, k):
        return math.nan if r[k] is None else r[k]
    ax2.plot(x, [_f(r, "freq_low") for r in w], label="low or zero", color="#2b6cb0", lw=1.6)
    ax2.plot(x, [_f(r, "freq_zero") for r in w], label="zero", color="#222222", lw=0.8, ls="--")
    ax2.plot(x, [_f(r, "freq_high") for r in w], label="high", color="#c53030", lw=1)
    ax2.set_ylim(-0.02, 1.02)
    ax2.set_ylabel("fraction of samples")
    ax2.set_xlabel(f"window ({res['summary']['window']:,} bp)")
    # contig dividers — skipped for fragmented assemblies, where thousands of
    # lines cost most of the plotting time and memory and are unreadable anyway
    if len(res["layout"]) <= 200:
        pos = 0
        for _, _, n in res["layout"][:-1]:
            pos += n
            for a in (ax1, ax2):
                a.axvline(pos - 0.5, color="#888888", lw=0.5, ls=":")
    ax2.legend(fontsize=7, loc="upper right", ncol=3)
    ax1.set_title(f"{res['summary']['organism']}  —  {res['summary']['n_informative']} informative samples"
                  f" (* = below depth cutoff)", fontsize=9)
    fig.savefig(path, dpi=120)
    plt.close(fig)
    return True


def main(argv=None):
    args = parse_args(argv)
    refs, n_samples = load(args)
    print(f"[alignment_trends] {n_samples} sample(s), {len(refs)} reference(s) with depth profiles")
    summaries, windows, regions, samples, matrix, results = [], [], [], [], [], []
    for key, ref in refs.items():
        res = analyse_reference(ref, args)
        if not res:
            continue
        samples.extend(res["samples"])
        if res.get("skipped"):
            continue
        summaries.append(res["summary"])
        windows.extend(res["windows"])
        regions.extend(res["regions"])
        matrix.extend(res["matrix"])
        results.append(res)
    summaries.sort(key=lambda s: (-s["n_informative"], s["organism"]))
    regions.sort(key=lambda r: (r["organism"], r["type"], r["contig"], r["start"]))

    pre = args.output
    _write_tsv(f"{pre}.regions.tsv", regions, REGION_COLS)
    _write_tsv(f"{pre}.windows.tsv", windows, WINDOW_COLS)
    _write_tsv(f"{pre}.samples.tsv", samples, SAMPLE_COLS)
    _write_tsv(f"{pre}.summary.tsv", summaries, SUMMARY_COLS)
    if args.matrix:
        _write_tsv(f"{pre}.matrix.tsv", matrix)
    params = {k: v for k, v in vars(args).items() if k not in ("input",)}
    with open(f"{pre}.json", "w") as fh:
        # The per-window table stays in PREFIX.windows.tsv only: it is by far
        # the largest piece and the report recomputes windows from the JSONs.
        json.dump({"taxtriage_alignment_trends": True, "version": 1, "params": params,
                   "n_samples": n_samples, "references": summaries, "regions": regions,
                   "samples": samples}, fh, separators=(",", ":"))
    if _write_xlsx(f"{pre}.xlsx", [("summary", SUMMARY_COLS, summaries), ("regions", REGION_COLS, regions),
                                   ("samples", SAMPLE_COLS, samples), ("windows", WINDOW_COLS, windows)]):
        print(f"[alignment_trends] wrote {pre}.xlsx")
    if args.plots > 0 and results:
        os.makedirs(f"{pre}.plots", exist_ok=True)
        top = sorted(results, key=lambda r: (-r["summary"]["n_informative"], -r["summary"]["total_len"]))
        made = 0
        for r in top[:args.plots]:
            safe = "".join(c if c.isalnum() or c in "-_." else "_" for c in r["summary"]["organism"])[:80]
            if _plot(r, os.path.join(f"{pre}.plots", f"{r['summary']['key']}_{safe}.png")):
                made += 1
        print(f"[alignment_trends] wrote {made} plot(s) to {pre}.plots/")
    print(f"[alignment_trends] {len(summaries)} reference(s) analysed; "
          f"{sum(1 for r in regions if r['type'] == 'zero')} zero, "
          f"{sum(1 for r in regions if r['type'] == 'low')} low, "
          f"{sum(1 for r in regions if r['type'] == 'high')} high recurrent region(s)")


if __name__ == "__main__":
    main()
